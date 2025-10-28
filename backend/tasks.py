import celery_app
import models

celery_app = celery_app.celery_app
from sqlalchemy.orm import sessionmaker
from sqlalchemy import create_engine, and_, func
import datetime
import os
from dotenv import load_dotenv

load_dotenv()

# Database setup
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./test.db")  # Use SQLite for development
engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

@celery_app.task
def extract_text_from_file(file_id: str):
    """Extract text from uploaded file and update database"""
    db = SessionLocal()
    try:
        file_obj = db.query(models.File).filter(models.File.id == file_id).first()
        if not file_obj:
            return {"error": "File not found"}

        # Read file content from storage (assuming local storage for now)
        file_path = file_obj.file_path
        if not os.path.exists(file_path):
            return {"error": "File not found on disk"}

        with open(file_path, 'rb') as f:
            file_content = f.read()

        # Extract text based on file type
        extracted_text = _extract_text_from_content(file_content, file_obj.filename)

        if extracted_text:
            file_obj.extracted_text = extracted_text
            db.commit()

            # Generate embeddings and store in vector DB and embeddings_index
            try:
                from ai_engine import ai_engine
                if ai_engine.embedding_model and ai_engine.vector_db:
                    embedding = ai_engine.embedding_model.encode(extracted_text).tolist()

                    # Store in vector database
                    ai_engine.vector_db.upsert([(
                        file_id,
                        embedding,
                        {
                            "text": extracted_text,
                            "source": f"user_upload_{file_obj.user_id}",
                            "filename": file_obj.filename,
                            "user_id": file_obj.user_id,
                            "type": "file_content"
                        }
                    )])

                    # Store in embeddings_index table
                    embedding_index = models.EmbeddingsIndex(
                        file_id=file_id,
                        embedding_vector=embedding,
                        metadata_json={
                            "filename": file_obj.filename,
                            "mime_type": file_obj.mime_type,
                            "file_size": file_obj.file_size
                        }
                    )
                    db.add(embedding_index)
                    db.commit()

            except Exception as embed_error:
                logger.error(f"Error generating embeddings for file {file_id}: {embed_error}")
                # Don't fail the task if embedding fails

            return {"status": "success", "text_length": len(extracted_text)}
        else:
            return {"error": "Failed to extract text from file"}

    except Exception as e:
        db.rollback()
        logger.error(f"Error processing file {file_id}: {e}")
        return {"error": str(e)}
    finally:
        db.close()

def _extract_text_from_content(file_content: bytes, filename: str) -> str:
    """Extract text from file content based on file type"""
    try:
        file_extension = filename.lower().split('.')[-1]

        if file_extension == 'pdf':
            # PDF extraction
            try:
                import PyPDF2
                pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
                return text
            except ImportError:
                return "PDF processing not available"

        elif file_extension in ['docx', 'doc']:
            # Word document extraction
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_content))
                text = ""
                for paragraph in doc.paragraphs:
                    text += paragraph.text + "\n"
                return text
            except ImportError:
                return "Word document processing not available"

        elif file_extension in ['xlsx', 'xls']:
            # Excel extraction
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_content))
                text = ""
                for sheet in wb:
                    for row in sheet.iter_rows(values_only=True):
                        text += " ".join([str(cell) for cell in row if cell]) + "\n"
                return text
            except ImportError:
                return "Excel processing not available"

        elif file_extension in ['png', 'jpg', 'jpeg']:
            # OCR for images
            try:
                import pytesseract
                from PIL import Image
                image = Image.open(io.BytesIO(file_content))
                text = pytesseract.image_to_string(image)
                return text
            except ImportError:
                return "OCR processing not available"

        else:
            # Plain text or unknown format
            return file_content.decode('utf-8', errors='ignore')

    except Exception as e:
        logger.error(f"Error extracting text from {filename}: {e}")
        return ""

@celery_app.task
def delete_old_chats():
    """Delete AI chats older than 7 days"""
    db = SessionLocal()
    try:
        cutoff_date = datetime.datetime.utcnow() - datetime.timedelta(days=7)

        # Delete old chats
        deleted_count = db.query(models.AI_Chat).filter(
            models.AI_Chat.created_at < cutoff_date
        ).delete()

        db.commit()
        return {"status": "success", "deleted_chats": deleted_count}
    except Exception as e:
        db.rollback()
        return {"error": str(e)}
    finally:
        db.close()

@celery_app.task
def grant_daily_free_coins():
    """Grant 4 free coins to users who haven't received them today"""
    db = SessionLocal()
    try:
        today = datetime.datetime.utcnow().date()

        # Find users who haven't received free coins today
        users_to_grant = db.query(models.User).filter(
            ~models.User.id.in_(
                db.query(models.Coins.user_id).filter(
                    func.date(models.Coins.free_daily_coins_last_grant_date) == today
                )
            )
        ).all()

        granted_count = 0
        for user in users_to_grant:
            # Get or create coins record
            coins_record = db.query(models.Coins).filter(models.Coins.user_id == user.id).first()
            if not coins_record:
                coins_record = models.Coins(user_id=user.id, balance=0)
                db.add(coins_record)

            # Grant 4 coins
            coins_record.balance += 4
            coins_record.free_daily_coins_last_grant_date = datetime.datetime.utcnow()

            # Create transaction record
            transaction = models.Transaction(
                user_id=user.id,
                type="free_grant",
                amount=4,
                description="Daily free coins"
            )
            db.add(transaction)
            granted_count += 1

        db.commit()
        return {"status": "success", "granted_to_users": granted_count}
    except Exception as e:
        db.rollback()
        return {"error": str(e)}
    finally:
        db.close()

@celery_app.task
def retrain_ai_model():
    """Re-training/curation pipeline task (manual approval)"""
    # TODO: Implement AI model re-training logic
    # This would involve:
    # 1. Collecting new training data
    # 2. Preprocessing data
    # 3. Training the model
    # 4. Validating performance
    # 5. Deploying the new model
    return {"status": "success", "message": "Re-training pipeline initiated"}
